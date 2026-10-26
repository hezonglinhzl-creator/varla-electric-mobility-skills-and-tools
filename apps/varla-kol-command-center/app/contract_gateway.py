"""Preview-only boundary for the installed KOL contract updater."""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
from pathlib import Path
import shutil
import subprocess
from typing import Any, Callable
from uuid import uuid4

from docx import Document

from app.gmail_gateway import sanitize_error


COMMAND_TIMEOUT_SECONDS = 30
Runner = Callable[..., Any]


@dataclass(frozen=True, slots=True)
class ContractPreview:
    output_path: Path
    source_sha256: str
    output_sha256: str
    structure_unchanged: bool


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _structure(path: Path) -> tuple[object, ...]:
    document = Document(path)
    paragraphs = tuple(
        paragraph.text
        for paragraph in document.paragraphs
        if not paragraph.text.strip().lower().startswith(
            "impact id number:"
        )
    )
    table_shapes = tuple(
        tuple(len(row.cells) for row in table.rows)
        for table in document.tables
    )
    return (
        len(paragraphs),
        len(document.tables),
        table_shapes,
        len(document.sections),
    )


class ContractGateway:
    """Create an isolated contract preview; no send action is exposed."""

    def __init__(
        self,
        *,
        python_executable: str | Path,
        script_path: str | Path,
        output_root: str | Path,
        runner: Runner | None = None,
        timeout_seconds: int = COMMAND_TIMEOUT_SECONDS,
    ) -> None:
        self.python_executable = str(python_executable)
        self.script_path = Path(script_path)
        self.output_root = Path(output_root)
        self.runner = runner or subprocess.run
        self.timeout_seconds = timeout_seconds

    def preview(
        self,
        source_path: str | Path,
        *,
        name: str,
        email: str,
        rate: str,
        agreement_date: str,
        bike_model: str | None = None,
        retain_output: bool = True,
    ) -> ContractPreview:
        source = Path(source_path).resolve()
        output_root = self.output_root.resolve()
        if output_root == source or (
            output_root.exists() and not output_root.is_dir()
        ):
            raise ValueError("contract output target must be a directory")
        if not source.is_file():
            raise FileNotFoundError("contract source document is unavailable")
        if source.suffix.lower() != ".docx":
            raise ValueError("contract source must be a DOCX file")

        before_hash = _sha256(source)
        before_structure = _structure(source)
        output_root.mkdir(parents=True, exist_ok=True)
        output = (
            output_root
            / f"contract-preview-{uuid4().hex}.docx"
        ).resolve()
        isolated_input = (
            output_root / f".contract-input-{uuid4().hex}.docx"
        ).resolve()
        if output == source:
            raise ValueError("contract preview cannot overwrite its source")

        argv = [
            self.python_executable,
            str(self.script_path),
            "--input",
            str(isolated_input),
            "--output",
            str(output),
            "--name",
            name,
            "--email",
            email,
            "--rate",
            rate,
            "--date",
            agreement_date,
        ]
        if bike_model:
            argv.extend(["--bike-model", bike_model])

        succeeded = False
        try:
            try:
                shutil.copyfile(source, isolated_input)
            except OSError as exc:
                raise RuntimeError(sanitize_error(exc)) from exc
            try:
                completed = self.runner(
                    argv,
                    capture_output=True,
                    text=True,
                    timeout=self.timeout_seconds,
                    check=False,
                    shell=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise RuntimeError("contract preview timed out") from exc
            except OSError as exc:
                raise RuntimeError(sanitize_error(exc)) from exc

            if getattr(completed, "returncode", 1) != 0:
                detail = sanitize_error(
                    getattr(completed, "stderr", "")
                    or getattr(completed, "stdout", "")
                )
                raise RuntimeError(f"contract preview failed: {detail}")
            if not output.is_file():
                raise RuntimeError("contract updater did not create a preview")
            after_structure = _structure(output)
            if after_structure != before_structure:
                raise RuntimeError(
                    "contract preview changed document structure"
                )
            result = ContractPreview(
                output_path=output,
                source_sha256=before_hash,
                output_sha256=_sha256(output),
                structure_unchanged=True,
            )
            succeeded = True
            return result
        finally:
            try:
                source_intact = _sha256(source) == before_hash
            except OSError:
                source_intact = False
            isolated_input.unlink(missing_ok=True)
            if output.is_file() and (
                not succeeded or not retain_output or not source_intact
            ):
                output.unlink()
            if not source_intact:
                raise RuntimeError("contract source changed during preview")


__all__ = ["ContractGateway", "ContractPreview"]
