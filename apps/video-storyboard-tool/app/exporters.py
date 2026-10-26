"""Export storyboard rows to CSV, Markdown, and Excel with optional images."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .models import FEISHU_FIELD_MAP, FEISHU_TABLE_FIELDS, Storyboard, TABLE_FIELDS


def export_csv(storyboard: Storyboard, output_path: Path) -> Path:
    rows = storyboard.to_table_rows()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=TABLE_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in TABLE_FIELDS})
    return output_path


def export_markdown(storyboard: Storyboard, output_path: Path) -> Path:
    bundle = export_feishu_bundle(storyboard, output_path.parent, output_path.stem)
    return bundle["storyboard_md"]


def clean_cell(value: object) -> str:
    text = str(value or "").replace("\r\n", " ").replace("\n", " ").replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace("|", "｜")


def feishu_payload_rows(storyboard: Storyboard) -> list[dict[str, str]]:
    payload = storyboard.to_feishu_payload()
    return [{key: clean_cell(row.get(key, "")) for key in FEISHU_FIELD_MAP} for row in payload["rows"]]


def feishu_markdown_table(storyboard: Storyboard) -> str:
    rows = feishu_payload_rows(storyboard)
    lines = [
        "| " + " | ".join(FEISHU_TABLE_FIELDS) + " |",
        "|---|---|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            "| "
            + " | ".join(
                [
                    row["time"],
                    row["visual_reference"],
                    row["scene_description"],
                    row["action_camera"],
                    row["dialogue_sound"],
                    row["notes"],
                ]
            )
            + " |"
        )
    return "\n".join(lines)


def feishu_tsv_table(storyboard: Storyboard) -> str:
    rows = feishu_payload_rows(storyboard)
    lines = ["\t".join(FEISHU_TABLE_FIELDS)]
    for row in rows:
        lines.append(
            "\t".join(
                [
                    row["time"],
                    row["visual_reference"],
                    row["scene_description"],
                    row["action_camera"],
                    row["dialogue_sound"],
                    row["notes"],
                ]
            )
        )
    return "\n".join(lines) + "\n"


def export_feishu_bundle(storyboard: Storyboard, output_dir: Path, stem: str = "storyboard") -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = storyboard.to_feishu_payload()
    payload["rows"] = feishu_payload_rows(storyboard)

    storyboard_md = output_dir / f"{stem}.md"
    feishu_md = output_dir / "storyboard_feishu_table.md"
    feishu_tsv = output_dir / "storyboard_feishu_table.tsv"
    payload_json = output_dir / "storyboard_payload.json"

    table = feishu_markdown_table(storyboard)
    storyboard_md.write_text(
        "\n".join(
            [
                "# 视频拆解分镜脚本",
                "",
                "## 视频信息",
                "",
                f"- 视频名称：{storyboard.video_info.file_name}",
                f"- 视频时长：{storyboard.video_info.duration}",
                f"- 视频比例：{storyboard.video_info.aspect_ratio}",
                f"- 分辨率：{storyboard.video_info.resolution}",
                "",
                "## 飞书原生表格",
                "",
                table,
                "",
            ]
        ),
        encoding="utf-8",
    )
    feishu_md.write_text(table + "\n", encoding="utf-8")
    feishu_tsv.write_text(feishu_tsv_table(storyboard), encoding="utf-8-sig")
    payload_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {
        "storyboard_md": storyboard_md,
        "feishu_md": feishu_md,
        "feishu_tsv": feishu_tsv,
        "payload_json": payload_json,
    }


def export_legacy_markdown(storyboard: Storyboard, output_path: Path) -> Path:
    lines = [
        "# 视频拆解分镜脚本",
        "",
        "## 视频信息",
        "",
        f"- 视频名称：{storyboard.video_info.file_name}",
        f"- 视频时长：{storyboard.video_info.duration}",
        f"- 视频比例：{storyboard.video_info.aspect_ratio}",
        f"- 分辨率：{storyboard.video_info.resolution}",
        "",
        "## 分镜表",
        "",
        "| 镜号 | 景别 | 拍摄手法 | 画面 | 画面说明 | 备注 | 台词 | 英文台词 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for shot in storyboard.shots:
        frame = shot.frame_path.replace("\\", "/")
        lines.append(
            "| "
            + " | ".join(
                [
                    shot.shot_id,
                    shot.shot_size,
                    shot.camera_method,
                    frame,
                    shot.visual_description,
                    shot.notes,
                    shot.dialogue_cn,
                    shot.dialogue_en,
                ]
            ).replace("\n", "<br>")
            + " |"
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return output_path


def export_excel(storyboard: Storyboard, output_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "分镜脚本"
    ws.append(TABLE_FIELDS)

    for col_index, field in enumerate(TABLE_FIELDS, start=1):
        cell = ws.cell(row=1, column=col_index)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_index)].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 34
    ws.column_dimensions["F"].width = 24
    ws.column_dimensions["G"].width = 30
    ws.column_dimensions["H"].width = 30

    for row_index, shot in enumerate(storyboard.shots, start=2):
        values = [
            shot.shot_id,
            shot.shot_size,
            shot.camera_method,
            shot.frame_path,
            shot.visual_description,
            shot.notes,
            shot.dialogue_cn,
            shot.dialogue_en,
        ]
        ws.append(values)
        for col_index in range(1, len(TABLE_FIELDS) + 1):
            ws.cell(row=row_index, column=col_index).alignment = Alignment(vertical="top", wrap_text=True)
        frame = Path(shot.frame_path)
        if frame.exists():
            try:
                image = XLImage(str(frame))
                fit_image(image, max_width=110, max_height=196)
                ws.row_dimensions[row_index].height = max(118, image.height * 0.75 + 8)
                ws.add_image(image, f"D{row_index}")
                ws.cell(row=row_index, column=4).value = ""
            except Exception:
                ws.cell(row=row_index, column=4).value = str(frame)
        else:
            ws.row_dimensions[row_index].height = 92

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def export_word(storyboard: Storyboard, output_path: Path) -> Path:
    return export_word_batch([storyboard], output_path)


def export_word_batch(storyboards: list[Storyboard], output_path: Path) -> Path:
    """Create a Word document with native tables and embedded frame images."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)

    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(8)

    for index, storyboard in enumerate(storyboards):
        if index:
            document.add_page_break()
        heading = document.add_heading(storyboard.video_info.file_name or "分镜表", level=1)
        heading.paragraph_format.space_after = Pt(4)
        info = document.add_paragraph(
            f"{storyboard.video_info.duration} · {storyboard.video_info.aspect_ratio} · "
            f"{storyboard.video_info.resolution} · {storyboard.product_model or '未指定产品'}"
        )
        info.paragraph_format.space_after = Pt(8)
        add_storyboard_word_table(document, storyboard)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def add_storyboard_word_table(document: Document, storyboard: Storyboard) -> None:
    table = document.add_table(rows=1, cols=len(TABLE_FIELDS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    widths = [0.45, 0.65, 0.8, 1.15, 2.15, 1.65, 1.6, 1.6]
    header_cells = table.rows[0].cells
    for cell, field, width in zip(header_cells, TABLE_FIELDS, widths):
        cell.text = field
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8)

    for shot in storyboard.shots:
        row = table.add_row()
        row.height = Inches(1.45)
        values = [
            shot.shot_id,
            shot.shot_size,
            shot.camera_method,
            "",
            shot.visual_description,
            shot.notes,
            shot.dialogue_cn,
            shot.dialogue_en,
        ]
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.text = clean_cell(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)

        frame = Path(shot.frame_path)
        image_cell = row.cells[3]
        image_cell.text = ""
        if frame.exists():
            paragraph = image_cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            try:
                run.add_picture(str(frame), width=Inches(0.95))
            except Exception:
                image_cell.text = str(frame)
        else:
            image_cell.text = ""


def fit_image(image: XLImage, max_width: int, max_height: int) -> None:
    """Fit an Excel image into a box without changing its aspect ratio."""
    original_width = float(image.width or max_width)
    original_height = float(image.height or max_height)
    scale = min(max_width / original_width, max_height / original_height)
    image.width = int(original_width * scale)
    image.height = int(original_height * scale)
